"""
SIGNALFIRE AI - Pitch Deck Generator
Creates an editable PowerPoint presentation (pitch_deck.pptx) and PDF version (pitch_deck.pdf)
with professional light-theme enterprise typography, cards, tables, and accents.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

def generate_pptx(output_path: str = "deliverables/pitch_deck.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    NAVY = RGBColor(15, 23, 42)      # #0f172a
    LAVENDER = RGBColor(99, 102, 241) # #6366f1
    TEXT_MUTED = RGBColor(100, 116, 139) # #64748b
    BG_CARD = RGBColor(248, 250, 252)
    BORDER_CARD = RGBColor(226, 232, 240)
    WHITE = RGBColor(255, 255, 255)
    RED_ACCENT = RGBColor(239, 68, 68)
    GREEN_ACCENT = RGBColor(16, 185, 129)

    slides_content = [
        # Slide 1: Title
        {
            "title": "SIGNALFIRE AI",
            "subtitle": "Find the warning signs hidden across your data.",
            "category": "AI-Powered Business Intelligence, Anomaly Detection & Cross-Dataset Risk Discovery",
            "bullets": [
                "Demonstration Domain: Enterprise Retail & Supply Chain Operations",
                "Built with Exasol Personal, FastAPI, React, and Robust Multi-Metric Analytics",
                "Autonomous Hackathon Submission 2026"
            ]
        },
        # Slide 2: Problem Statement
        {
            "title": "The Operational Problem",
            "subtitle": "Why siloed metrics obscure developing operational crises",
            "text": (
                "Retail operations teams often struggle to identify emerging operational problems because sales, "
                "inventory, delivery, refund, and customer complaint data are analyzed separately. "
                "Small changes across individual metrics may appear harmless, even when their combined pattern "
                "indicates a developing problem. Manually reviewing these datasets is time-consuming and depends "
                "heavily on predefined alerts. SIGNALFIRE AI aims to help analysts discover unusual cross-metric patterns, "
                "investigate supporting evidence, and prioritize potential operational risks before they escalate."
            ),
            "bullets": [
                "Target Audience: Retail operations managers, supply chain analysts, and risk teams.",
                "The Impact: Uncoordinated alerts result in delayed root-cause diagnosis and compounding customer churn.",
                "The Boundary: We do not claim all monitoring systems fail—only that siloed metric alerts miss multi-table cascades."
            ]
        },
        # Slide 3: Why Existing Dashboards Miss the Pattern
        {
            "title": "Why Existing Dashboards Miss Cross-Metric Risks",
            "subtitle": "Independent metric thresholds fail when small changes co-occur across tables",
            "bullets": [
                "Independent Alerting: Delivery alerts only fire on severe transit breaches, while refunds and complaints have separate arbitrary cutoffs.",
                "Lack of Temporal Synchronization: A delivery delay today creates customer complaints in 2 days and returns in 4 days. Unaligned tools miss the lag.",
                "Excessive False Positives: Static percentage thresholds fire erratically on normal weekend seasonality without noise-robust MAD filtering.",
                "Absence of Unaffected Control Cohorts: Dashboards lack automated contrast between localized incident regions (North) and unaffected controls (South, East, West)."
            ]
        },
        # Slide 4: The SIGNALFIRE AI Solution
        {
            "title": "The SIGNALFIRE AI Solution",
            "subtitle": "A principled 5-stage analytical pipeline that separates facts from hypotheses",
            "bullets": [
                "1. Independent Anomaly Detection: Robust Z-scores via Median Absolute Deviation (MAD) and Interquartile Range (IQR) fences across 8 relational tables.",
                "2. Multi-Dataset Temporal Alignment: Daily/weekly synchronization isolating co-occurring signals across orders, dispatches, complaints, and returns.",
                "3. Statistical Association Engine: Empirical Pearson and Spearman correlation analysis with automated lead-lag offset estimation.",
                "4. Unaffected Cohort Benchmarking: Automated contrast of target incident segments against non-incident control cohorts.",
                "5. Disciplined Synthesis: Explicit, tamper-proof separation of Observed Facts, Statistical Associations, Hypotheses, Competing Explanations, and Recommendations."
            ]
        },
        # Slide 5: Product Walkthrough & User Journey
        {
            "title": "Live Product Architecture & User Journey",
            "subtitle": "From raw table ingestion to decision-support and audit-ready reporting",
            "bullets": [
                "Overview Dashboard: Live KPI pulse, signal onset trend curves, risk severity breakdown, and live database engine indicator.",
                "Data Explorer & Profiler: Schema inspector, missing value percentages, record preview with search/filter, and safe CSV upload.",
                "Signal Detection Engine: Granular inspection of observed vs baseline metrics, confidence scores, and detection methods.",
                "Cross-Metric Investigation: Co-occurrence clusters, lead-lag offsets, and control cohort comparisons.",
                "Interactive Risk Map & Workspace: Topological segment network, priority assignments (P1-P3), and conclusion logging.",
                "What-If Stress Testing & Export: Arithmetic vs model projections, and one-click Executive PDF/JSON/CSV downloads."
            ]
        },
        # Slide 6: Technical Architecture & Exasol Integration
        {
            "title": "Technical Architecture & Exasol Integration",
            "subtitle": "High-throughput SQL analytics with seamless local fallback transparency",
            "bullets": [
                "Analytical Storage Layer: Native Exasol Personal integration via official pyexasol driver (WebSocket JSON-RPC).",
                "Database Adapter Architecture: Unified DatabaseAdapter interface supporting Exasol Personal and DuckDB in-process OLAP.",
                "Status Transparency: Clear indicator in UI and API ('Exasol Personal' vs 'DuckDB Local'). Never mislabels local calculations.",
                "API & Analytical Core: FastAPI, Pydantic v2, Pandas, NumPy, Scikit-learn, SciPy, and ReportLab.",
                "Frontend Design Language: React, TypeScript, Vite, Tailwind CSS, Recharts, and refined 3D glassmorphism."
            ]
        },
        # Slide 7: Innovation & Honest Differentiation
        {
            "title": "Innovation & Differentiating Principles",
            "subtitle": "Credible, research-backed capabilities without hype or false claims",
            "bullets": [
                "Non-Conflation of Correlation & Causation: The system strictly presents statistical associations as hypotheses and systematically surfaces competing explanations.",
                "Robust Outlier Normalization: Uses Median Absolute Deviation (MAD), eliminating sensitivity to extreme noise that compromises standard deviations.",
                "Automated Control Cohort Validation: Validates that regional shifts are genuine localized anomalies rather than system-wide macro trends.",
                "Evidence-Backed AI Assistant: Operates via deterministic analytical tool retrieval. Never hallucinates numbers or executes unrestricted SQL."
            ]
        },
        # Slide 8: Analytical Validation & Ground Truth Results
        {
            "title": "Analytical Evaluation & Empirical Metrics",
            "subtitle": "Evaluated strictly against recorded ground-truth synthetic anomaly injections",
            "bullets": [
                "SignalFire Multi-Metric Detector: Precision = 71.4% | Recall = 100.0% | F1-Score = 0.833 (0 False Negatives across all injected crises).",
                "Simple %-Change Baseline Detector: Precision = 45.5% | Recall = 100.0% | F1-Score = 0.625 (Suffers 6 False Positives from seasonal volume noise).",
                "Detection Delay: Anomalies isolated within 1 to 3 observation windows of onset.",
                "Limitation Disclosure: Ground truth is established on synthetic distributions; empirical validation on production non-stationary data is required."
            ]
        },
        # Slide 9: Target Users & Operational Impact
        {
            "title": "Target Users & Operational Decision Support",
            "subtitle": "Empowering cross-functional teams to act before operational issues compound",
            "bullets": [
                "Retail Operations Managers: Uncover localized carrier hub breakdowns before customer delivery complaints trigger mass returns.",
                "Supply Chain Analysts: Trace inventory stockouts directly to upstream transit delays and supplier defect rates.",
                "Customer Experience & Retention Teams: Understand why repeat purchase rates decline by connecting fulfillment failures to churn.",
                "Risk & Performance Executives: Receive executive PDF summaries with audit trails and verified conclusions for SLA negotiations."
            ]
        },
        # Slide 10: Future Roadmap & Closing
        {
            "title": "Future Roadmap & Project Conclusion",
            "subtitle": "Scaling cross-dataset intelligence for enterprise operations",
            "bullets": [
                "Distributed Exasol Analytics: Scaling from Exasol Personal to multi-node Exasol clusters for multi-billion record datasets.",
                "Unstructured Data Fusion: Ingesting customer support chat transcripts and courier tracking telemetry via multimodal LLM embeddings.",
                "Automated Workflow Triggers: Direct webhook webhooks to ERP systems (SAP, NetSuite) for automated carrier SLA penalty issuance.",
                "SIGNALFIRE AI: Find the warning signs hidden across your data."
            ]
        }
    ]

    for slide_data in slides_content:
        slide = prs.slides.add_slide(blank_layout)

        # Header Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.1))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = slide_data["title"]
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = NAVY

        p_sub = tf.add_paragraph()
        p_sub.text = slide_data["subtitle"]
        p_sub.font.size = Pt(13)
        p_sub.font.color.rgb = LAVENDER

        # Content Card
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.8), Inches(1.8), Inches(11.733), Inches(5.1)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = BG_CARD
        shape.line.color.rgb = BORDER_CARD
        shape.line.width = Pt(1.5)

        tb_content = slide.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(11.1), Inches(4.7))
        tf_c = tb_content.text_frame
        tf_c.word_wrap = True

        if "text" in slide_data:
            p_text = tf_c.paragraphs[0]
            p_text.text = slide_data["text"]
            p_text.font.size = Pt(13)
            p_text.font.color.rgb = NAVY
            p_text.space_after = Pt(14)
            first_bullet = True
        else:
            first_bullet = False

        for idx, bullet in enumerate(slide_data.get("bullets", [])):
            if not first_bullet and idx == 0 and "text" not in slide_data:
                p_b = tf_c.paragraphs[0]
            else:
                p_b = tf_c.add_paragraph()
            p_b.text = f"• {bullet}"
            p_b.font.size = Pt(13.5)
            p_b.font.color.rgb = NAVY
            p_b.space_after = Pt(10)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    prs.save(output_path)
    print(f"PowerPoint Pitch Deck saved to {output_path}")

def generate_pdf(output_path: str = "deliverables/pitch_deck.pdf"):
    doc = SimpleDocTemplate(
        output_path,
        pagesize=landscape(letter),
        rightMargin=40,
        leftMargin=40,
        topMargin=35,
        bottomMargin=35
    )

    styles = getSampleStyleSheet()
    primary = colors.HexColor("#0f172a")
    accent = colors.HexColor("#6366f1")
    light_bg = colors.HexColor("#f8fafc")
    border_color = colors.HexColor("#cbd5e1")

    title_style = ParagraphStyle('SlideTitle', fontSize=22, leading=26, textColor=primary, fontName='Helvetica-Bold')
    sub_style = ParagraphStyle('SlideSub', fontSize=11, leading=15, textColor=accent, fontName='Helvetica')
    body_style = ParagraphStyle('SlideBody', fontSize=10.5, leading=16, textColor=primary, fontName='Helvetica')

    story = []

    slides = [
        ("Slide 1: SIGNALFIRE AI", "Find the warning signs hidden across your data", [
            "<b>Category:</b> AI-Powered Business Intelligence, Anomaly Detection & Cross-Dataset Risk Discovery",
            "<b>Domain:</b> Enterprise Retail Operations & Supply Chain Telemetry",
            "<b>Built with:</b> Exasol Personal, DuckDB Fallback, FastAPI, React, and Robust Multi-Metric Analytics",
            "<b>Mission:</b> Help operations analysts discover unusual cross-dataset patterns and prioritize emerging risks before they escalate."
        ]),
        ("Slide 2: The Operational Problem", "Why siloed operational metrics obscure emerging systemic failures", [
            "<b>Siloed Monitoring:</b> Sales, deliveries, complaints, refunds, and retention are analyzed separately.",
            "<b>Developing Crises:</b> Small, sub-threshold changes across individual metrics co-occur to form severe failures.",
            "<b>High Review Burden:</b> Manual review is time-consuming and heavily dependent on uncoordinated single-metric alerts.",
            "<b>Target Audience:</b> Business operations managers, supply chain analysts, and performance leaders."
        ]),
        ("Slide 3: Why Existing Dashboards Miss Cross-Metric Risks", "Independent thresholds fail when small changes co-occur across tables", [
            "<b>Independent Alerting:</b> Delivery alerts fire only on large transit delays, missing early warning signals in complaints.",
            "<b>Temporal Lag Disconnection:</b> Transit delays lead return requests by 3-4 days; unaligned dashboards miss the lead-lag offset.",
            "<b>Seasonality Noise:</b> Static percentage thresholds generate false alarms during normal weekend volume surges.",
            "<b>No Control Cohorts:</b> Lacks automated contrast against unaffected operational segments (e.g. North vs South/East/West)."
        ]),
        ("Slide 4: The SIGNALFIRE AI Solution", "A principled 5-stage analytical pipeline separating facts from hypotheses", [
            "<b>1. Independent Anomaly Detection:</b> Robust Z-Scores via Median Absolute Deviation (MAD) and IQR fences.",
            "<b>2. Multi-Dataset Temporal Alignment:</b> Synchronizing 8 tables onto unified 24h temporal windows.",
            "<b>3. Statistical Association Engine:</b> Empirical Pearson and Spearman correlation with lead-lag offset estimation.",
            "<b>4. Unaffected Cohort Benchmarking:</b> Automated contrast against non-incident regional controls.",
            "<b>5. Disciplined Synthesis:</b> Explicit separation of Facts, Statistical Associations, Hypotheses, and Limitations."
        ]),
        ("Slide 5: Product Walkthrough & Demonstration", "A unified decision-support workspace for analysts and leaders", [
            "<b>Overview Dashboard:</b> Real-time KPI cards, weekly signal trend curves, risk breakdown, and active engine badge.",
            "<b>Data Explorer & Profiler:</b> Schema inspector, null distributions, paginated record preview, and safe CSV upload.",
            "<b>Signal Detection Engine:</b> Complete evidence trails, observed vs baseline values, and detection methodologies.",
            "<b>Cross-Metric Engine:</b> Prioritized risk clusters, empirical correlation metrics, and control cohort benchmarks.",
            "<b>Interactive Risk Map & Workspace:</b> Topological segment graph, priority assignments (P1-P3), and conclusion logging.",
            "<b>What-If Stress Testing & Export:</b> Arithmetic vs model projections, and one-click Executive PDF/JSON/CSV downloads."
        ]),
        ("Slide 6: Technical Architecture & Exasol Integration", "High-throughput in-memory SQL analytics with transparent local fallback", [
            "<b>Analytical Storage:</b> Exasol Personal integration via official pyexasol driver (WebSocket JSON-RPC).",
            "<b>Adapter Architecture:</b> Abstract DatabaseAdapter supporting Exasol Personal and DuckDB OLAP fallback.",
            "<b>Full Transparency:</b> Live engine status displayed across UI and API. Never mislabels local calculations.",
            "<b>Backend Engine:</b> Python 3.12, FastAPI, Pydantic v2, Pandas, NumPy, Scikit-learn, SciPy, and ReportLab.",
            "<b>Frontend Experience:</b> React 18, TypeScript, Vite, Tailwind CSS, Recharts, and 3D glassmorphism design language."
        ]),
        ("Slide 7: Innovation & Differentiating Principles", "Credible, research-backed capabilities without hype or false claims", [
            "<b>Correlation != Causation:</b> System never claims correlation proves causation. Explicitly formulates competing hypotheses.",
            "<b>Robust MAD Normalization:</b> Median Absolute Deviation suppresses extreme noise that distorts standard deviations.",
            "<b>Automated Control Cohort Validation:</b> Verifies localized shocks against unaffected regional cohorts.",
            "<b>Evidence-Backed AI Assistant:</b> Tool-grounded query architecture that executes verified SQL analytics."
        ]),
        ("Slide 8: Analytical Evaluation & Empirical Metrics", "Evaluated strictly against recorded ground-truth synthetic anomaly injections", [
            "<b>SignalFire Multi-Metric Detector:</b> Precision = 71.4% | Recall = 100.0% | F1-Score = 0.833 (Zero false negatives).",
            "<b>Simple %-Change Baseline Detector:</b> Precision = 45.5% | Recall = 100.0% | F1-Score = 0.625 (High false positive rate).",
            "<b>Detection Latency:</b> Signals isolated within 1 to 3 observation windows of onset.",
            "<b>Honest Limitations:</b> Ground truth established on synthetic data; non-stationary production distributions require ongoing calibration."
        ]),
        ("Slide 9: Target Users & Operational Impact", "Empowering cross-functional teams to act before operational issues compound", [
            "<b>Retail Operations Managers:</b> Detect carrier bottlenecks before transit delays trigger customer return waves.",
            "<b>Supply Chain Analysts:</b> Connect warehouse stockouts directly to upstream supplier defect and transit delays.",
            "<b>Customer Experience Teams:</b> Mitigate repeat purchase churn by identifying fulfillment root causes.",
            "<b>Leadership & Executives:</b> Receive audit-ready PDF reports with verified conclusions for SLA partner negotiations."
        ]),
        ("Slide 10: Future Roadmap & Closing", "Scaling cross-dataset intelligence for enterprise operations", [
            "<b>Distributed Exasol Clusters:</b> Scaling from Exasol Personal to multi-node clusters for billion-record datasets.",
            "<b>Multimodal Telemetry Fusion:</b> Incorporating unstructured customer support tickets and dispatch logs.",
            "<b>Automated ERP Orchestration:</b> Triggering automated supplier notifications and SLA penalty workflows in SAP/NetSuite.",
            "<b>SIGNALFIRE AI:</b> Find the warning signs hidden across your data."
        ])
    ]

    for title, subtitle, bullets in slides:
        story.append(Paragraph(title, title_style))
        story.append(Paragraph(subtitle, sub_style))
        story.append(HRFlowable(width="100%", thickness=1, color=accent, spaceBefore=4, spaceAfter=10))

        bullet_paragraphs = [
            Paragraph(f"• {b}", body_style) for b in bullets
        ]
        t = Table([[bp] for bp in bullet_paragraphs], colWidths=[700])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), light_bg),
            ('BOX', (0, 0), (-1, -1), 0.5, border_color),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('LEFTPADDING', (0, 0), (-1, -1), 12),
            ('RIGHTPADDING', (0, 0), (-1, -1), 12),
        ]))
        story.append(t)
        story.append(Spacer(1, 40))

    doc.build(story)
    print(f"PDF Pitch Deck saved to {output_path}")

if __name__ == "__main__":
    generate_pptx("deliverables/pitch_deck.pptx")
    generate_pdf("deliverables/pitch_deck.pdf")
